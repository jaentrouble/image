import openpyxl as xl

class Saver() :
    def __init__(self) :
        try :
            self.wb = xl.load_workbook(filename = 'savefile.xlsx')
        except FileNotFoundError :
            self.wb = xl.Workbook()
            self.wb.save('savefile.xlsx')
        self.ws = self.wb.worksheets[0]
        self.rec_col = 1
        self.img_col = 2
        self.row = 1

    def save(self, record_list : list, img_idx : list):
        try :
            while self.ws.cell(self.row,self.rec_col).value != None :
                self.row += 1
            for idx in range(len(record_list)) :
                self.ws.cell(self.row, self.rec_col).value = record_list[idx]
                self.ws.cell(self.row, self.img_col).value = img_idx[idx]
                self.row += 1
                self.wb.save('savefile.xlsx')
        except :
            return False
        return True